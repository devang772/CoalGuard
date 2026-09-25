"""Monthly compliance report: collect the numbers, then write them as PDF (fpdf2) and Excel (openpyxl)."""
import hashlib
import io
from collections import Counter, defaultdict
from datetime import datetime, timedelta

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.config import settings
from app.models import (Attendance, Capa, ComplianceTask, Contractor, EnvReading, Finding, Inspection, Obligation,
                        Observation, OrgUnit, ProductionLog)
from app.services.dashboard import SUSPICIOUS_DISPATCH_RATIO, month_bounds
from app.services.fraud import contractor_alerts, contractor_score
from app.services.tasks import compliance_stats
from app.utils import IST_OFFSET, ist_date, ist_day_start_utc, today_ist

TITLE = "Monthly Compliance Report"


# ---------------------------------------------------------------- data

def collect(db: Session, mine_ids: list[int], month: str, scope_label: str, generated_by: str,
            generated_at: datetime) -> dict:
    first, last = month_bounds(month)
    last = min(last, today_ist())
    start, end = ist_day_start_utc(first), ist_day_start_utc(last + timedelta(days=1))
    mines = {m.id: m.name for m in db.scalars(select(OrgUnit).where(OrgUnit.id.in_(mine_ids)))}

    compliance = compliance_stats(db, mine_ids, first, last)
    per_obligation: dict[tuple, Counter] = defaultdict(Counter)
    for code, title, law_ref, freq, t_status, due, done_at in db.execute(
            select(Obligation.code, Obligation.title, Obligation.law_ref, Obligation.frequency, ComplianceTask.status,
                   ComplianceTask.due_date, ComplianceTask.done_at)
            .join(Obligation, Obligation.id == ComplianceTask.obligation_id)
            .where(ComplianceTask.mine_id.in_(mine_ids), ComplianceTask.due_date.between(first, last))):
        key = (code or "", title, law_ref, freq)
        per_obligation[key]["due"] += 1
        on_time = t_status == "done" and done_at is not None and ist_date(done_at) <= due
        per_obligation[key]["on_time" if on_time else "done_late" if t_status == "done" else "not_done"] += 1
    obligations = sorted(({"code": k[0], "title": k[1], "law_ref": k[2], "frequency": k[3],
                           **{field: v[field] for field in ("due", "on_time", "done_late", "not_done")}}
                          for k, v in per_obligation.items()), key=lambda o: (-o["not_done"], o["code"]))

    inspections = Counter(t for (t,) in db.execute(select(Inspection.type).where(
        Inspection.mine_id.in_(mine_ids), Inspection.started_at >= start, Inspection.started_at < end)))
    findings = Counter(s for (s,) in db.execute(select(Finding.severity).where(
        Finding.mine_id.in_(mine_ids), Finding.created_at >= start, Finding.created_at < end)))

    capas_opened = db.execute(select(Capa.status, Capa.escalation_level).where(
        Capa.mine_id.in_(mine_ids), Capa.created_at >= start, Capa.created_at < end)).all()
    capas = {"opened": len(capas_opened),
             "closed_in_month": db.scalar(select(func.count()).select_from(Capa).where(
                 Capa.mine_id.in_(mine_ids), Capa.closed_at >= start, Capa.closed_at < end)) or 0,
             "still_open_and_overdue": db.scalar(select(func.count()).select_from(Capa).where(
                 Capa.mine_id.in_(mine_ids), Capa.status.in_(["open", "rejected"]), Capa.due_at < end)) or 0,
             "escalated": sum(1 for _, level in capas_opened if level and level > 0)}

    obs_rows = db.execute(select(Observation.type, Observation.source, Observation.created_at, Observation.mine_id,
                                 Observation.text, Observation.severity)
                          .where(Observation.mine_id.in_(mine_ids), Observation.created_at >= start,
                                 Observation.created_at < end).order_by(Observation.created_at)).all()
    reports = Counter(t for t, *_ in obs_rows)
    incidents = [{"date": ist_date(c), "mine": mines.get(m, ""), "severity": sev, "text": text}
                 for t, _, c, m, text, sev in obs_rows if t == "incident"]

    att = db.execute(select(Attendance.valid, Attendance.gate_entry).where(
        Attendance.mine_id.in_(mine_ids), Attendance.time >= start, Attendance.time < end)).all()
    attendance = {"valid": sum(1 for v, _ in att if v), "invalid": sum(1 for v, _ in att if not v),
                  "without_gate_entry": sum(1 for v, g in att if v and not g)}

    contractors = list(db.scalars(select(Contractor).where(Contractor.mine_id.in_(mine_ids))))
    alerts = contractor_alerts(db, [c.id for c in contractors])
    contractor_rows = sorted(({"name": c.name, "mine": mines.get(c.mine_id, ""), "score": contractor_score(alerts[c.id]),
                               "alerts": "; ".join(a["title"] for a in alerts[c.id]) or "-"} for c in contractors),
                             key=lambda r: r["score"])

    env = defaultdict(list)
    for mine_id, pm10 in db.execute(select(EnvReading.mine_id, EnvReading.pm10).where(
            EnvReading.mine_id.in_(mine_ids), EnvReading.time >= start, EnvReading.time < end)):
        if pm10 is not None:
            env[mine_id].append(pm10)
    environment = [{"mine": mines[m], "avg_pm10": round(sum(v) / len(v), 1), "max_pm10": round(max(v), 1),
                    "days_over_limit": sum(1 for x in v if x > settings.pm10_limit)} for m, v in sorted(env.items())]

    prod = defaultdict(lambda: {"produced": 0.0, "dispatched": 0.0, "gap_days": []})
    for mine_id, d, produced, dispatched in db.execute(select(ProductionLog.mine_id, ProductionLog.date,
                                                              ProductionLog.produced_t, ProductionLog.dispatched_t)
                                                       .where(ProductionLog.mine_id.in_(mine_ids),
                                                              ProductionLog.date.between(first, last))):
        prod[mine_id]["produced"] += produced
        prod[mine_id]["dispatched"] += dispatched
        if produced and dispatched / produced < SUSPICIOUS_DISPATCH_RATIO:
            prod[mine_id]["gap_days"].append(d.strftime("%d %b"))
    production = [{"mine": mines[m], "produced_t": round(v["produced"]), "dispatched_t": round(v["dispatched"]),
                   "suspicious_days": ", ".join(v["gap_days"]) or "-"} for m, v in sorted(prod.items())]

    return {"title": TITLE, "scope": scope_label, "month": month, "period": f"{first:%d %b %Y} - {last:%d %b %Y}",
            "mines": sorted(mines.values()), "generated_by": generated_by,
            "generated_at": (generated_at + IST_OFFSET).strftime("%d %b %Y %I:%M %p IST"),
            "compliance": compliance, "obligations": obligations, "inspections": dict(inspections),
            "findings": dict(findings), "capas": capas, "reports": dict(reports), "incidents": incidents,
            "attendance": attendance, "contractors": contractor_rows, "environment": environment,
            "pm10_limit": settings.pm10_limit, "production": production}


# ---------------------------------------------------------------- PDF

def _latin(text) -> str:
    """Core PDF fonts only have Latin-1 characters."""
    value = str(text).replace("₹", "Rs ").replace("–", "-").replace("—", "-").replace("…", "...")
    return value.encode("latin-1", "replace").decode("latin-1")


class _Pdf(FPDF):
    def __init__(self, data: dict):
        super().__init__(orientation="P", unit="mm", format="A4")
        self.data = data
        self.set_auto_page_break(auto=True, margin=18)
        self.set_margins(14, 14, 14)

    def header(self):
        self.set_fill_color(30, 41, 59)
        self.rect(0, 0, 210, 11, "F")
        self.set_font("Helvetica", "B", 9)
        self.set_text_color(255, 255, 255)
        self.set_xy(14, 3)
        self.cell(0, 5, _latin(f"KHANAN NETRA  |  {self.data['title']}  |  {self.data['scope']}  |  {self.data['month']}"))
        self.set_text_color(30, 41, 59)
        self.set_y(16)

    def footer(self):
        self.set_y(-12)
        self.set_font("Helvetica", "", 7)
        self.set_text_color(100, 116, 139)
        self.cell(0, 5, _latin(f"Generated {self.data['generated_at']} by {self.data['generated_by']}  |  "
                               f"Verify this file's SHA-256 fingerprint in Khanan Netra (Reports > Verify)  |  "
                               f"Page {self.page_no()}"), align="C")

    def heading(self, text: str):
        self.ln(3)
        self.set_font("Helvetica", "B", 12)
        self.set_text_color(30, 41, 59)
        self.cell(0, 7, _latin(text), new_x="LMARGIN", new_y="NEXT")
        self.set_draw_color(245, 158, 11)
        self.line(14, self.get_y(), 196, self.get_y())
        self.ln(2)

    def pairs(self, rows: list[tuple[str, object]]):
        self.set_font("Helvetica", "", 9)
        for label, value in rows:
            self.set_text_color(71, 85, 105)
            self.cell(70, 6, _latin(label))
            self.set_text_color(15, 23, 42)
            self.cell(0, 6, _latin(value if value is not None else "-"), new_x="LMARGIN", new_y="NEXT")

    def table(self, headers: list[str], rows: list[list], widths: list[int]):
        if not rows:
            self.set_font("Helvetica", "I", 9)
            self.cell(0, 6, "No records in this period.", new_x="LMARGIN", new_y="NEXT")
            return
        self.set_font("Helvetica", "B", 8)
        self.set_fill_color(241, 245, 249)
        for h, w in zip(headers, widths):
            self.cell(w, 6, _latin(h), border=1, fill=True)
        self.ln()
        self.set_font("Helvetica", "", 8)
        for row in rows:
            if self.get_y() > 270:
                self.add_page()
            for value, w in zip(row, widths):
                text = _latin(value if value is not None else "-")
                while text and self.get_string_width(text) > w - 2:
                    text = text[:-2]
                self.cell(w, 6, text, border=1)
            self.ln()


def to_pdf(data: dict) -> bytes:
    pdf = _Pdf(data)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(0, 10, _latin(data["title"]), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    pdf.cell(0, 6, _latin(f"{data['scope']}  -  {data['period']}"), new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 6, _latin(f"Mines: {', '.join(data['mines'])}"), new_x="LMARGIN", new_y="NEXT")

    c = data["compliance"]
    pdf.heading("1. Compliance summary")
    pdf.pairs([("Compliance (done on time / due)", f"{c['compliance_pct']}%" if c["compliance_pct"] is not None else "-"),
               ("Tasks due", c["due"]), ("Done on time", c["done_on_time"]), ("Done late", c["done_late"]),
               ("Not done (overdue)", c["overdue"])])
    pdf.table(["Category", "Due", "On time", "Compliance %"],
              [[x["category"], x["due"], x["done_on_time"], x["compliance_pct"]] for x in c["by_category"]],
              [60, 30, 30, 40])

    pdf.heading("2. Obligations (most missed first)")
    pdf.table(["Code", "Obligation", "Law", "Freq", "Due", "On time", "Late", "Missed"],
              [[o["code"], o["title"], o["law_ref"], o["frequency"], o["due"], o["on_time"], o["done_late"],
                o["not_done"]] for o in data["obligations"][:40]],
              [24, 56, 36, 15, 11, 14, 11, 15])

    pdf.heading("3. Inspections and findings")
    pdf.pairs([("Inspections by type", ", ".join(f"{k}: {v}" for k, v in sorted(data["inspections"].items())) or "0"),
               ("Findings by severity", ", ".join(f"{k}: {v}" for k, v in sorted(data["findings"].items())) or "0")])

    cp = data["capas"]
    pdf.heading("4. Corrective actions (CAPA)")
    pdf.pairs([("Opened this month", cp["opened"]), ("Closed this month", cp["closed_in_month"]),
               ("Escalated (opened this month)", cp["escalated"]),
               ("Still open and overdue at period end", cp["still_open_and_overdue"])])

    pdf.heading("5. Field reports and incidents")
    pdf.pairs([("Reports by type", ", ".join(f"{k.replace('_', ' ')}: {v}" for k, v in sorted(data["reports"].items())) or "0")])
    pdf.table(["Date", "Mine", "Severity", "Incident"],
              [[i["date"].strftime("%d %b"), i["mine"], i["severity"], i["text"]] for i in data["incidents"][:30]],
              [18, 38, 20, 106])

    a = data["attendance"]
    pdf.heading("6. Attendance and contractors")
    pdf.pairs([("Valid attendance records", a["valid"]), ("Refused attempts", a["invalid"]),
               ("Valid but without gate entry", a["without_gate_entry"])])
    pdf.table(["Contractor", "Mine", "Score", "Alerts (current)"],
              [[r["name"], r["mine"], r["score"], r["alerts"]] for r in data["contractors"][:25]],
              [48, 34, 14, 86])

    pdf.heading(f"7. Environment (PM10 limit {data['pm10_limit']:g} ug/m3)")
    pdf.table(["Mine", "Average PM10", "Highest PM10", "Days over limit"],
              [[e["mine"], e["avg_pm10"], e["max_pm10"], e["days_over_limit"]] for e in data["environment"]],
              [60, 40, 40, 42])

    pdf.heading("8. Production vs dispatch")
    pdf.table(["Mine", "Produced (t)", "Dispatched (t)", f"Days dispatched < {int(SUSPICIOUS_DISPATCH_RATIO * 100)}%"],
              [[p["mine"], f"{p['produced_t']:,}", f"{p['dispatched_t']:,}", p["suspicious_days"]]
               for p in data["production"]], [50, 35, 35, 62])
    return bytes(pdf.output())


# ---------------------------------------------------------------- Excel

def to_xlsx(data: dict) -> bytes:
    wb = Workbook()
    bold, head_fill = Font(bold=True), PatternFill("solid", fgColor="1E293B")

    def sheet(title: str, headers: list[str], rows: list[list], first: bool = False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(headers)
        for cell in ws[1]:
            cell.font, cell.fill = Font(bold=True, color="FFFFFF"), head_fill
        for row in rows:
            ws.append(row)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(60, max(10, *(len(str(c.value or "")) + 2 for c in col)))
        ws.freeze_panes = "A2"
        return ws

    c = data["compliance"]
    ws = sheet("Summary", ["Item", "Value"], [
        ["Report", data["title"]], ["Scope", data["scope"]], ["Period", data["period"]],
        ["Mines", ", ".join(data["mines"])], ["Generated", f"{data['generated_at']} by {data['generated_by']}"],
        ["Compliance % (done on time / due)", c["compliance_pct"]], ["Tasks due", c["due"]],
        ["Done on time", c["done_on_time"]], ["Done late", c["done_late"]], ["Not done (overdue)", c["overdue"]],
        ["CAPAs opened", data["capas"]["opened"]], ["CAPAs closed", data["capas"]["closed_in_month"]],
        ["CAPAs escalated", data["capas"]["escalated"]],
        ["CAPAs open and overdue at period end", data["capas"]["still_open_and_overdue"]],
        ["Valid attendance", data["attendance"]["valid"]], ["Refused attendance", data["attendance"]["invalid"]],
        ["Valid without gate entry", data["attendance"]["without_gate_entry"]]], first=True)
    for cell in ws["A"]:
        cell.font = bold
    sheet("Compliance by category", ["Category", "Due", "Done on time", "Compliance %"],
          [[x["category"], x["due"], x["done_on_time"], x["compliance_pct"]] for x in c["by_category"]])
    sheet("Obligations", ["Code", "Obligation", "Law reference", "Frequency", "Due", "On time", "Late", "Missed"],
          [[o["code"], o["title"], o["law_ref"], o["frequency"], o["due"], o["on_time"], o["done_late"], o["not_done"]]
           for o in data["obligations"]])
    sheet("Inspections & findings", ["Type", "Count"],
          [[f"Inspection: {k}", v] for k, v in sorted(data["inspections"].items())]
          + [[f"Finding: {k}", v] for k, v in sorted(data["findings"].items())])
    sheet("Field reports", ["Type", "Count"], [[k, v] for k, v in sorted(data["reports"].items())])
    ws = sheet("Incidents", ["Date", "Mine", "Severity", "Description"],
               [[i["date"], i["mine"], i["severity"], i["text"]] for i in data["incidents"]])
    for cell in ws["D"]:
        cell.alignment = Alignment(wrap_text=True)
    sheet("Contractors", ["Contractor", "Mine", "Score", "Alerts (current)"],
          [[r["name"], r["mine"], r["score"], r["alerts"]] for r in data["contractors"]])
    sheet("Environment", ["Mine", "Average PM10", "Highest PM10", f"Days over {data['pm10_limit']:g}"],
          [[e["mine"], e["avg_pm10"], e["max_pm10"], e["days_over_limit"]] for e in data["environment"]])
    sheet("Production", ["Mine", "Produced (t)", "Dispatched (t)", "Suspicious dispatch days"],
          [[p["mine"], p["produced_t"], p["dispatched_t"], p["suspicious_days"]] for p in data["production"]])
    out = io.BytesIO()
    wb.properties.title = f"{data['title']} {data['scope']} {data['month']}"
    wb.save(out)
    return out.getvalue()


def fingerprint(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()
