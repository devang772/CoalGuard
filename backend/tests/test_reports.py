"""Module 9: command dashboard, mine dashboard, leaderboard, PDF/Excel reports with fingerprint + approval."""
import io

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.config import settings
from app.db import SessionLocal
from app.models import Capa, ComplianceTask, OrgUnit, User
from app.services import storage
from app.utils import today_ist

CIL, BCCL, GM, MANAGER, REGULATOR, CONTRACTOR, OFFICER, SUPERVISOR, WORKER = (
    f"90000000{i:02d}" for i in range(1, 10))
MONTH = today_ist().strftime("%Y-%m")


@pytest.fixture(scope="module", autouse=True)
def _data(seeded_data):
    return seeded_data


def mine_id(name):
    with SessionLocal() as db:
        return db.scalar(select(OrgUnit.id).where(OrgUnit.name == name))


# ---------------------------------------------------------------- dashboards

def test_command_summary_matches_the_database(client, login):
    s = client.get("/dashboard/summary", headers=login(CIL)).json()
    with SessionLocal() as db:
        overdue = db.scalar(select(func.count()).select_from(ComplianceTask).where(ComplianceTask.status == "overdue"))
        open_capas = db.scalar(select(func.count()).select_from(Capa).where(
            Capa.status.in_(["open", "in_review", "rejected"])))
    assert s["scope"]["mines"] == 12 and s["overdue_tasks"] == overdue
    assert s["open_capas"]["total"] == open_capas
    assert s["open_capas"]["lt7"] + s["open_capas"]["d7_30"] + s["open_capas"]["gt30"] == open_capas
    assert 0 <= s["compliance_pct"] <= 100 and len(s["compliance_trend"]) == 6 and len(s["incidents_trend"]) == 6
    assert len(s["top_risky_mines"]) == 5 and s["top_risky_mines"][0]["risk_pct"] >= s["top_risky_mines"][-1]["risk_pct"]
    assert set(s) >= {"incidents_month", "near_miss_month", "avg_trust_score", "active_workers_today",
                      "invalid_attendance_today", "capa_by_category", "recent_alerts", "active_sos"}


def test_summary_is_scoped(client, login):
    assert client.get("/dashboard/summary", headers=login(GM)).json()["scope"]["mines"] == 2
    assert client.get("/dashboard/summary", params={"org_id": mine_id("Kusunda OCP")},
                      headers=login(CIL)).json()["scope"]["mines"] == 1
    assert client.get("/dashboard/summary", headers=login(WORKER)).status_code == 403


def test_mine_dashboard_shows_planted_problems(client, login):
    basta = client.get(f"/dashboard/mine/{mine_id('Bastacolla OCP')}", headers=login(GM)).json()
    assert basta["suspicious_dispatch_days"] >= 4
    assert all(p["dispatch_ratio"] < 0.8 for p in basta["production"] if p["suspicious"])
    ashoka = client.get(f"/dashboard/mine/{mine_id('Ashoka OCP')}", headers=login(CIL)).json()
    assert ashoka["pm10_days_over_limit"] >= 4 and ashoka["pm10_limit"] == settings.pm10_limit
    assert basta["mine"]["name"] == "Bastacolla OCP" and "compliance_30d" in basta
    assert client.get(f"/dashboard/mine/{mine_id('Ashoka OCP')}", headers=login(GM)).status_code == 403


def test_leaderboard(client, login):
    board = client.get("/dashboard/leaderboard", params={"month": MONTH}, headers=login(CIL)).json()
    rows = board["rows"]
    assert len(rows) == 12 and [r["rank"] for r in rows] == list(range(1, 13))
    assert rows[-1]["name"] == "Kusunda OCP"
    assert all(0 <= r["safety_score"] <= 100 and r["trend"] in ("up", "down", "same") for r in rows)
    assert set(rows[0]["breakdown"]) == {"compliance", "capa_on_time", "photo_trust", "no_incidents", "overdue_penalty"}
    subs = client.get("/dashboard/leaderboard", params={"month": MONTH, "by": "subsidiary"}, headers=login(CIL)).json()
    assert {r["code"] for r in subs["rows"]} == {"BCCL", "CCL", "MCL"}
    assert client.get("/dashboard/leaderboard", params={"month": "2026-13"}, headers=login(CIL)).status_code == 422


# ---------------------------------------------------------------- reports

@pytest.fixture(scope="module")
def moonidih_reports(client, login):
    r = client.post("/reports", json={"month": MONTH, "mine_id": mine_id("Moonidih UG")}, headers=login(MANAGER))
    assert r.status_code == 201, r.text
    return r.json()


def test_generate_pdf_and_excel(client, moonidih_reports):
    reports = {x["format"]: x for x in moonidih_reports["reports"]}
    assert set(reports) == {"pdf", "xlsx"}
    assert all(x["status"] == "generated" and x["stored_in"] == "local" and len(x["sha256"]) == 64
               for x in reports.values())
    pdf = client.get(reports["pdf"]["url"])                       # signed link, no header
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF") and len(pdf.content) == reports["pdf"]["size_bytes"]
    assert "attachment" in pdf.headers["content-disposition"]
    xlsx = client.get(reports["xlsx"]["url"])
    wb = load_workbook(io.BytesIO(xlsx.content))
    assert wb.sheetnames[:3] == ["Summary", "Compliance by category", "Obligations"]
    summary = {row[0].value: row[1].value for row in wb["Summary"].iter_rows(min_row=2)}
    assert summary["Scope"] == "Moonidih UG" and summary["Report"] == "Monthly Compliance Report"
    assert moonidih_reports["summary"]["tasks_due"] >= 0


def test_verify_detects_changes(client, login, moonidih_reports):
    pdf = next(x for x in moonidih_reports["reports"] if x["format"] == "pdf")
    original = client.get(pdf["url"]).content
    ok = client.post("/reports/verify", files={"file": ("r.pdf", original, "application/pdf")}, headers=login(WORKER)).json()
    assert ok["match"] is True and ok["report"]["id"] == pdf["id"]
    middle = len(original) // 2
    tampered = original[:middle] + bytes([original[middle] ^ 1]) + original[middle + 1:]      # one byte changed
    bad = client.post("/reports/verify", files={"file": ("r.pdf", tampered, "application/pdf")}, headers=login(WORKER)).json()
    assert bad["match"] is False and "changed" in bad["message"]


def test_report_approval_two_person_rule(client, login, moonidih_reports):
    pdf = next(x for x in moonidih_reports["reports"] if x["format"] == "pdf")
    body = {"entity": "report", "entity_id": pdf["id"], "decision": "approve"}
    own = client.post("/approvals", json=body, headers=login(MANAGER))
    assert own.status_code == 403 and "you generated this report" in own.json()["detail"]
    assert client.post("/approvals", json=body, headers=login(GM)).status_code == 201
    after = client.get(f"/reports/{pdf['id']}", headers=login(MANAGER)).json()
    assert after["status"] == "approved" and after["approved_by_name"].startswith("Rajesh")
    assert after["approvals_verified"] is True
    assert client.post("/approvals", json=body, headers=login(BCCL)).status_code == 409
    history = client.get("/approvals", params={"entity": "report", "entity_id": pdf["id"]}, headers=login(GM)).json()
    assert [h["decision"] for h in history] == ["approve"]


def test_area_report_and_access(client, login, moonidih_reports):
    area = client.post("/reports", json={"month": MONTH, "formats": ["pdf"]}, headers=login(GM))
    assert area.status_code == 201 and area.json()["reports"][0]["scope_label"] == "Jharia Area"
    assert client.post("/reports", json={"month": MONTH}, headers=login(WORKER)).status_code == 403
    assert client.post("/reports", json={"month": MONTH, "mine_id": mine_id("Kusunda OCP")},
                       headers=login(MANAGER)).status_code == 403
    assert client.post("/reports", json={"month": MONTH, "formats": ["pdf"], "mine_id": mine_id("Moonidih UG")},
                       headers=login(REGULATOR)).status_code == 201
    history = client.get("/reports", params={"month": MONTH}, headers=login(GM)).json()
    assert history["total"] >= 4 and {r["scope_label"] for r in history["items"]} <= {"Moonidih UG", "Jharia Area"}
    with SessionLocal() as db:
        kusunda_manager = db.scalar(select(User.phone).where(User.name == "Mine Manager, Kusunda OCP"))
    pdf_id = next(x["id"] for x in moonidih_reports["reports"] if x["format"] == "pdf")
    assert client.get(f"/reports/{pdf_id}", headers=login(kusunda_manager)).status_code == 403
    assert client.get("/reports", headers=login(kusunda_manager)).json()["total"] == 0
    assert client.get(f"/reports/{pdf_id}/file").status_code == 401


def test_reports_can_go_to_cloudinary(client, login, monkeypatch):
    stored = {}
    monkeypatch.setattr(settings, "storage_backend", "cloudinary")
    monkeypatch.setattr(settings, "cloudinary_url", "cloudinary://1:secret@demo-cloud")
    monkeypatch.setattr(storage, "_cloud_upload", lambda data, public_id, resource_type:
                        stored.setdefault("x", {"public_id": public_id, "resource_type": resource_type}))
    monkeypatch.setattr(storage, "_cloud_download_url", lambda public_id, fmt, resource_type, expires_at:
                        f"https://api.cloudinary.com/v1_1/demo-cloud/{resource_type}/download?public_id={public_id}")
    r = client.post("/reports", json={"month": MONTH, "mine_id": mine_id("Moonidih UG"), "formats": ["pdf"]},
                    headers=login(MANAGER)).json()["reports"][0]
    assert r["stored_in"] == "cloudinary" and stored["x"]["resource_type"] == "raw"
    assert stored["x"]["public_id"].startswith("khanan-netra/reports/") and stored["x"]["public_id"].endswith(".pdf")
    redirect = client.get(r["url"], follow_redirects=False)
    assert redirect.status_code == 307 and "/raw/download" in redirect.headers["location"]
